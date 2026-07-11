"""
Databricks Delta Live Tables pipeline — signal documentation ingestion
=========================================================================

Extracts text from PDF/DOCX/PPTX/XLSX signal documentation files (e.g. OEM
CAN matrix / interface control documents, SOME/IP service specs, glossaries)
uploaded to a Unity Catalog Volume, into a single Delta table for Databricks
Genie Space to search when grounding natural-language questions about signal
meaning.

Separate from blf_ingestion (dlt_blf_pipeline.py): this pipeline shares no
lineage with blf_gold_signals and doesn't need the vector_blf wheel. It's
triggered (not continuous) since documentation uploads are infrequent, unlike
the streaming BLF telemetry ingestion.

Layer layout
------------
  blf_signal_doc_sections    table — text extracted from PDF/DOCX/PPTX/XLSX signal docs

Setup
-----
Set pipeline parameters (Edit -> Advanced -> Parameters):
    blf.signal_docs_path        /Volumes/mycat/myschema/signals/docs   (required)
    blf.semantic_model_endpoint databricks-claude-3-7-sonnet           (optional)

blf.signal_docs_path points at a Volume directory of PDF/DOCX/PPTX/XLSX
signal documentation files. Auto Loader extracts text per PDF page, PPTX
slide, XLSX sheet, or whole DOCX document into blf_signal_doc_sections.

blf.semantic_model_endpoint, if set, names a Model Serving endpoint used via
the ai_query() SQL function to derive a semantic_summary column from each
section's text; the pipeline's run-as identity needs CAN_QUERY on that
endpoint. Left unset, semantic_summary stays NULL and no model calls happen.

This pipeline is triggered: run it manually (`databricks bundle run
signal_docs`) or on a schedule after uploading new documentation files.
"""

from __future__ import annotations

import dlt
import pandas as pd
import pyspark.sql.functions as F
from pyspark.sql.types import (
    IntegerType,
    LongType,
    StringType,
    StructField,
    StructType,
    TimestampType,
)

# ── pipeline parameters ───────────────────────────────────────────────────────

SIGNAL_DOCS_PATH = spark.conf.get("blf.signal_docs_path")
SEMANTIC_MODEL_ENDPOINT = spark.conf.get("blf.semantic_model_endpoint", "")

# ── signal documentation output schema (flat rows emitted by _parse_doc_batch) ─

_SIGNAL_DOC_SECTIONS_SCHEMA = StructType(
    [
        StructField("_source_file", StringType(), nullable=False),
        StructField("_file_mtime", TimestampType()),
        StructField("_file_size_bytes", LongType()),
        StructField("_ingested_at", TimestampType()),
        StructField("doc_type", StringType(), nullable=False),  # PDF | DOCX | PPTX | XLSX
        StructField("section_number", IntegerType(), nullable=False),
        StructField("section_label", StringType()),  # e.g. "Page 3", "Slide 5", "Sheet 'CAN Matrix'"
        StructField("text", StringType()),
    ]
)

# ── path helper ───────────────────────────────────────────────────────────────


def _local_path(spark_path: str) -> str:
    """Convert a Spark/DBFS URI to a local filesystem path.

    Unity Catalog Volumes are mounted at /Volumes/... without the dbfs: prefix.
    Legacy DBFS paths are mounted at /dbfs/....
    """
    if spark_path.startswith("dbfs:/Volumes/"):
        return spark_path[len("dbfs:") :]  # /Volumes/...
    if spark_path.startswith("dbfs:/"):
        return "/dbfs/" + spark_path[len("dbfs:/") :]
    return spark_path  # already a local path


# ── signal documentation extraction (PDF / DOCX / PPTX / XLSX) ────────────────


def _extract_pdf_sections(path: str) -> list[tuple[int, str, str]]:
    import pypdf

    reader = pypdf.PdfReader(path)
    return [(i, f"Page {i}", page.extract_text() or "") for i, page in enumerate(reader.pages, start=1)]


def _extract_docx_sections(path: str) -> list[tuple[int, str, str]]:
    import docx

    document = docx.Document(path)
    text = "\n".join(p.text for p in document.paragraphs if p.text)
    return [(1, "Document", text)]


def _extract_pptx_sections(path: str) -> list[tuple[int, str, str]]:
    from pptx import Presentation

    prs = Presentation(path)
    sections = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]
        sections.append((i, f"Slide {i}", "\n".join(t for t in texts if t)))
    return sections


def _extract_xlsx_sections(path: str) -> list[tuple[int, str, str]]:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sections = []
    for i, ws in enumerate(wb.worksheets, start=1):
        lines = ["\t".join(str(c) for c in row if c is not None) for row in ws.iter_rows(values_only=True)]
        sections.append((i, f"Sheet '{ws.title}'", "\n".join(line for line in lines if line)))
    return sections


_DOC_EXTRACTORS = {
    "pdf": ("PDF", _extract_pdf_sections),
    "docx": ("DOCX", _extract_docx_sections),
    "pptx": ("PPTX", _extract_pptx_sections),
    "xlsx": ("XLSX", _extract_xlsx_sections),
}


def _parse_doc_batch(iterator):
    """mapInPandas worker: extract per-section text from signal documentation
    files (PDF pages, DOCX whole-document, PPTX slides, XLSX sheets)."""
    from datetime import datetime, timezone

    for batch_df in iterator:
        now = datetime.now(timezone.utc)
        rows = []
        for _, row in batch_df.iterrows():
            spark_path = str(row["_source_file"])
            file_mtime = row.get("_file_mtime")
            file_size = row.get("_file_size_bytes")
            local = _local_path(spark_path)
            ext = local.rsplit(".", 1)[-1].lower()
            extractor = _DOC_EXTRACTORS.get(ext)
            if extractor is None:
                continue
            doc_type, extract_fn = extractor
            try:
                sections = extract_fn(local)
            except Exception as exc:
                print(f"[signal_docs_pipeline] failed to parse {spark_path!r}: {exc}")
                continue
            for section_number, section_label, text in sections:
                rows.append(
                    {
                        "_source_file": spark_path,
                        "_file_mtime": file_mtime,
                        "_file_size_bytes": file_size,
                        "_ingested_at": now,
                        "doc_type": doc_type,
                        "section_number": section_number,
                        "section_label": section_label,
                        "text": text,
                    }
                )
        if rows:
            yield pd.DataFrame(rows)


# ── signal documentation layer ─────────────────────────────────────────────────


@dlt.table(
    name="blf_signal_doc_sections",
    comment=(
        "Text sections extracted from signal documentation files (PDF "
        "pages, DOCX whole-document, PPTX slides, XLSX sheets) uploaded "
        "to the signals/docs Volume path. Search this table (e.g. WHERE "
        "text LIKE '%term%') to ground Genie Space questions about "
        "signal meaning that aren't obvious from signal_name alone in "
        "blf_gold_signals / blf_signal_catalog (produced by the separate "
        "blf_ingestion pipeline). semantic_summary is populated only when "
        "blf.semantic_model_endpoint is set; otherwise NULL."
    ),
    table_properties={
        "quality": "gold",
        "delta.autoOptimize.optimizeWrite": "true",
    },
)
def blf_signal_doc_sections():
    """Read signal documentation files via Auto Loader; extract text per
    page/slide/sheet/document, optionally enriched via ai_query()."""
    df = (
        spark.readStream.format("cloudFiles")
        .option("cloudFiles.format", "binaryFile")
        .option("pathGlobFilter", "*.{pdf,docx,pptx,xlsx}")
        .option(
            "cloudFiles.schemaLocation",
            f"{SIGNAL_DOCS_PATH}/_autoloader_schema",
        )
        .load(SIGNAL_DOCS_PATH)
        .select(
            F.col("path").alias("_source_file"),
            F.col("modificationTime").alias("_file_mtime"),
            F.col("length").alias("_file_size_bytes"),
        )
        .mapInPandas(_parse_doc_batch, schema=_SIGNAL_DOC_SECTIONS_SCHEMA)
    )
    if not SEMANTIC_MODEL_ENDPOINT:
        return df.withColumn("semantic_summary", F.lit(None).cast(StringType()))
    _endpoint_sql = SEMANTIC_MODEL_ENDPOINT.replace("'", "''")
    _prompt_prefix = (
        "Summarize this automotive signal documentation excerpt in 1-3 "
        "sentences. Call out any signal names, their physical meaning, "
        "and units if present. Excerpt: "
    )
    return df.withColumn(
        "semantic_summary",
        F.expr(f"ai_query('{_endpoint_sql}', concat('{_prompt_prefix}', text))"),
    )
