"""Tests for signal_docs_pipeline — pure-Python extraction helpers and mapInPandas worker."""

import importlib.util
from pathlib import Path

import pandas as pd
import pytest
import signal_docs_pipeline as pipeline

_needs_doc_libs = pytest.mark.skipif(
    not all(importlib.util.find_spec(mod) for mod in ("pypdf", "docx", "pptx", "openpyxl")),
    reason="pypdf/python-docx/python-pptx/openpyxl not installed; run: uv sync",
)


# ── _local_path tests ─────────────────────────────────────────────────────────


def test_local_path_dbfs_volume_prefix():
    assert pipeline._local_path("dbfs:/Volumes/main/blf_dev/docs/x.pdf") == "/Volumes/main/blf_dev/docs/x.pdf"


def test_local_path_legacy_dbfs():
    assert pipeline._local_path("dbfs:/tmp/x.pdf") == "/dbfs/tmp/x.pdf"


def test_local_path_already_local():
    assert pipeline._local_path("/Volumes/main/blf_dev/docs/x.pdf") == "/Volumes/main/blf_dev/docs/x.pdf"


def test_local_path_plain():
    assert pipeline._local_path("x.pdf") == "x.pdf"


# ── extraction helper tests ───────────────────────────────────────────────────


@_needs_doc_libs
def test_extract_pdf_sections(monkeypatch) -> None:
    class _FakePage:
        def __init__(self, text):
            self._text = text

        def extract_text(self):
            return self._text

    class _FakeReader:
        def __init__(self, path):
            self.pages = [_FakePage("Engine speed in rpm"), _FakePage("")]

    import pypdf

    monkeypatch.setattr(pypdf, "PdfReader", _FakeReader)
    sections = pipeline._extract_pdf_sections("dummy.pdf")
    assert sections == [(1, "Page 1", "Engine speed in rpm"), (2, "Page 2", "")]


@_needs_doc_libs
def test_extract_docx_sections(tmp_path: Path) -> None:
    import docx

    path = tmp_path / "doc.docx"
    document = docx.Document()
    document.add_paragraph("Engine speed in rpm")
    document.add_paragraph("")
    document.add_paragraph("Measured at the crankshaft position sensor.")
    document.save(str(path))

    sections = pipeline._extract_docx_sections(str(path))
    assert len(sections) == 1
    number, label, text = sections[0]
    assert number == 1
    assert label == "Document"
    assert "Engine speed in rpm" in text
    assert "crankshaft position sensor" in text


@_needs_doc_libs
def test_extract_pptx_sections(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    path = tmp_path / "doc.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    textbox.text_frame.text = "Engine speed in rpm"
    prs.save(str(path))

    sections = pipeline._extract_pptx_sections(str(path))
    assert len(sections) == 1
    number, label, text = sections[0]
    assert number == 1
    assert label == "Slide 1"
    assert "Engine speed in rpm" in text


@_needs_doc_libs
def test_extract_xlsx_sections(tmp_path: Path) -> None:
    import openpyxl

    path = tmp_path / "doc.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Signals"
    ws.append(["signal_name", "description"])
    ws.append(["EngineSpeed_rpm", "Engine speed in rpm"])
    wb.save(path)

    sections = pipeline._extract_xlsx_sections(str(path))
    assert len(sections) == 1
    number, label, text = sections[0]
    assert number == 1
    assert label == "Sheet 'Signals'"
    assert "EngineSpeed_rpm" in text
    assert "Engine speed in rpm" in text


# ── _parse_doc_batch tests ────────────────────────────────────────────────────


@_needs_doc_libs
def test_parse_doc_batch_dispatches_by_extension(tmp_path: Path) -> None:
    import openpyxl

    path = tmp_path / "doc.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["EngineSpeed_rpm", "Engine speed in rpm"])
    wb.save(path)

    batch_df = pd.DataFrame({"_source_file": [str(path)], "_file_mtime": [None], "_file_size_bytes": [None]})
    frames = list(pipeline._parse_doc_batch(iter([batch_df])))
    assert frames, "no rows parsed"
    df = pd.concat(frames, ignore_index=True)
    assert (df["doc_type"] == "XLSX").all()
    assert df.iloc[0]["section_number"] == 1
    assert "EngineSpeed_rpm" in df.iloc[0]["text"]


def test_parse_doc_batch_skips_unknown_extension() -> None:
    batch_df = pd.DataFrame({"_source_file": ["/tmp/foo.txt"], "_file_mtime": [None], "_file_size_bytes": [None]})
    frames = list(pipeline._parse_doc_batch(iter([batch_df])))
    assert frames == []
